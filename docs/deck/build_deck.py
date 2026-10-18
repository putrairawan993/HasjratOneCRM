"""Build the Hasjrat OneCRM assessment deck.

Usage (from the repository root):
    python3 docs/deck/build_deck.py

Writes docs/Hasjrat-OneCRM-Assessment.pptx. UI slides embed screenshots from
docs/screenshots/<name>.png when present and draw a labelled placeholder otherwise,
so re-run this script after capturing new screenshots.

PDF export (LibreOffice; in the Claude sandbox use the pptx skill's scripts/office/soffice.py wrapper):
    soffice --headless --convert-to pdf --outdir docs docs/Hasjrat-OneCRM-Assessment.pptx
"""
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls, qn
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'Hasjrat-OneCRM-Assessment.pptx'
SHOTS = ROOT / 'screenshots'

NAVY, GREEN, GREEN_DARK, GREEN_LIGHT = '152D36', '177F6C', '126653', 'E9F4EF'
INK, MUTED, LINE, BG, WHITE = '20343D', '5F6F74', 'E7EBE9', 'F6F8F6', 'FFFFFF'
AMBER, AMBER_LIGHT, RED, MINT, NAVY_SOFT = 'B58235', 'FBF3E6', 'B56550', 'A8D3B5', '24414C'
FONT = 'Calibri'
W, H = 13.333, 7.5
TOTAL = 23


def rgb(hex_color):
    return RGBColor.from_string(hex_color)


def new_slide(prs, bg=WHITE):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(bg)
    return slide


def text(slide, x, y, w, h, content, size=14, bold=False, color=INK, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, italic=False, line_spacing=None, name=None):
    """content: str, list of paragraphs (str or list of (text, opts) runs)."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        box.name = name
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    paragraphs = content if isinstance(content, list) else [content]
    for i, para in enumerate(paragraphs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if line_spacing:
            p.line_spacing = line_spacing
        runs = para if isinstance(para, list) else [(para, {})]
        for run_text, opts in runs:
            r = p.add_run()
            r.text = run_text
            f = r.font
            f.name = FONT
            f.size = Pt(opts.get('size', size))
            f.bold = opts.get('bold', bold)
            f.italic = opts.get('italic', italic)
            f.color.rgb = rgb(opts.get('color', color))
    return box


def bullets(slide, x, y, w, h, items, size=14, color=INK, space_after=6, bullet_color=GREEN):
    box = text(slide, x, y, w, h, items, size=size, color=color)
    for p in box.text_frame.paragraphs:
        pPr = p._p.get_or_add_pPr()
        pPr.set('marL', str(Inches(0.22)))
        pPr.set('indent', str(-Inches(0.22)))
        pPr.append(parse_xml(f'<a:buClr {nsdecls("a")}><a:srgbClr val="{bullet_color}"/></a:buClr>'))
        pPr.append(parse_xml(f'<a:buFont {nsdecls("a")} typeface="Arial"/>'))
        pPr.append(parse_xml(f'<a:buChar {nsdecls("a")} char="&#8226;"/>'))
        p.space_after = Pt(space_after)
    return box


def rect(slide, x, y, w, h, fill=GREEN_LIGHT, line=None, radius=0.08, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
         dash=False, line_width=1.0, name=None):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        s.name = name
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    if fill:
        s.fill.solid()
        s.fill.fore_color.rgb = rgb(fill)
    else:
        s.fill.background()
    if line:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(line_width)
        if dash:
            s.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    else:
        s.line.fill.background()
    style = s._element.find(qn('p:style'))
    if style is not None:
        s._element.remove(style)
    s.text_frame.text = ''
    return s


def badge(slide, x, y, d, label, fill=GREEN, color=WHITE, size=14):
    s = rect(slide, x, y, d, d, fill=fill, shape=MSO_SHAPE.OVAL)
    tf = s.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = label
    r.font.name, r.font.size, r.font.bold = FONT, Pt(size), True
    r.font.color.rgb = rgb(color)
    return s


def chip(slide, x, y, w, h, label, fill=GREEN_LIGHT, color=INK, size=12, bold=False, line=None,
         align=PP_ALIGN.CENTER):
    s = rect(slide, x, y, w, h, fill=fill, line=line, radius=0.25)
    tf = s.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.08)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    lines = label.split('\n')
    for i, line_text in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = line_text
        r.font.name, r.font.size = FONT, Pt(size if i == 0 else max(size - 2, 10))
        r.font.bold = bold if i == 0 else False
        r.font.color.rgb = rgb(color if i == 0 else (MUTED if fill in (WHITE, GREEN_LIGHT, BG, AMBER_LIGHT) else color))
    return s


def arrow(slide, x1, y1, x2, y2, color=GREEN, width=1.75):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(width)
    ln = c.line._get_or_add_ln()
    ln.append(parse_xml(f'<a:tailEnd {nsdecls("a")} type="triangle" w="med" len="med"/>'))
    return c


def line(slide, x1, y1, x2, y2, color=LINE, width=2.0):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(width)
    return c


def header(slide, eyebrow, title):
    text(slide, 0.6, 0.4, 12, 0.3, eyebrow.upper(), size=12, bold=True, color=GREEN)
    text(slide, 0.6, 0.7, 12.1, 0.75, title, size=30, bold=True, color=INK)


def footer(slide, n, dark=False):
    color = '8DA1A8' if dark else MUTED
    text(slide, 0.6, 7.02, 9, 0.25, 'Hasjrat OneCRM · Konsep assessment oleh Putra Irawan · data prototype sintetis',
         size=10, color=color)
    text(slide, 11.73, 7.02, 1.0, 0.25, f'{n} / {TOTAL}', size=10, color=color, align=PP_ALIGN.RIGHT)


def notes(slide, body):
    slide.notes_slide.notes_text_frame.text = body


def table(slide, x, y, w, rows, col_widths, size=12, header_size=12, row_h=0.42, header_fill=NAVY,
          first_col_bold=True, zebra=True):
    shape = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w),
                                   Inches(row_h * len(rows)))
    tbl = shape.table
    tbl.first_row = True
    tbl.horz_banding = False
    for i, cw in enumerate(col_widths):
        tbl.columns[i].width = Inches(cw)
    for r, row in enumerate(rows):
        tbl.rows[r].height = Inches(row_h)
        for c, value in enumerate(row):
            cell = tbl.cell(r, c)
            cell.margin_left = cell.margin_right = Inches(0.08)
            cell.margin_top = cell.margin_bottom = Inches(0.04)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            if r == 0:
                cell.fill.fore_color.rgb = rgb(header_fill)
            else:
                cell.fill.fore_color.rgb = rgb(BG if zebra and r % 2 == 0 else WHITE)
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            run = p.add_run()
            run.text = value
            run.font.name = FONT
            run.font.size = Pt(header_size if r == 0 else size)
            run.font.bold = r == 0 or (first_col_bold and c == 0)
            run.font.color.rgb = rgb(WHITE if r == 0 else INK)
    return shape


def picture(slide, name, x, y, w, h, label):
    path = SHOTS / f'{name}.png'
    if path.exists():
        with Image.open(path) as img:
            iw, ih = img.size
        scale = min(w / iw, h / ih)
        pw, ph = iw * scale, ih * scale
        px, py = x + (w - pw) / 2, y + (h - ph) / 2
        pic = slide.shapes.add_picture(str(path), Inches(px), Inches(py), Inches(pw), Inches(ph))
        pic.line.color.rgb = rgb(LINE)
        pic.line.width = Pt(1)
        pic.name = f'Screenshot {label}'
        return pic
    box = rect(slide, x, y, w, h, fill=BG, line='B8C4C7', dash=True, radius=0.04,
               name=f'Placeholder screenshot {label}')
    text(slide, x + 0.3, y + h / 2 - 0.45, w - 0.6, 0.4, f'Screenshot: {label}', size=16, bold=True, color=INK,
         align=PP_ALIGN.CENTER)
    text(slide, x + 0.3, y + h / 2 + 0.0, w - 0.6, 0.6,
         f'Simpan PNG di docs/screenshots/{name}.png lalu jalankan ulang docs/deck/build_deck.py',
         size=12, color=MUTED, align=PP_ALIGN.CENTER)
    return box


def card(slide, x, y, w, h, title, body, num=None, fill=GREEN_LIGHT, title_size=16, body_size=13,
         badge_fill=GREEN, body_color=INK):
    rect(slide, x, y, w, h, fill=fill)
    tx = x + 0.25
    if num is not None:
        badge(slide, x + 0.25, y + 0.25, 0.46, str(num), fill=badge_fill, size=15)
        tx = x + 0.88
    text(slide, tx, y + 0.27, w - (tx - x) - 0.2, 0.45, title, size=title_size, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    text(slide, x + 0.25, y + 0.9, w - 0.5, h - 1.05, body, size=body_size, color=body_color)


# ---------------------------------------------------------------- slides

def slide_title(prs):
    s = new_slide(prs, NAVY)
    mark = rect(s, 0.8, 0.75, 0.72, 0.76, fill=MINT, radius=0.18)
    tf = mark.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = 'H1'
    r.font.name, r.font.size, r.font.bold = FONT, Pt(22), True
    r.font.color.rgb = rgb('153C32')
    text(s, 1.7, 0.78, 4, 0.4, [[('hasjrat', {'bold': True, 'color': WHITE, 'size': 24}),
                                  ('.', {'bold': True, 'color': MINT, 'size': 24})]])
    text(s, 1.72, 1.2, 4, 0.3, 'ONECRM', size=11, bold=True, color='ADBEC2')
    text(s, 0.8, 2.25, 11.5, 0.4, 'TECHNICAL ASSESSMENT · SENIOR BUSINESS SYSTEM ANALYST & DIGITALIZATION',
         size=13, bold=True, color=MINT)
    text(s, 0.8, 2.75, 11.6, 1.7, 'Dari WhatsApp ke customer management platform yang terintegrasi',
         size=44, bold=True, color=WHITE, line_spacing=0.95)
    text(s, 0.8, 4.65, 11, 0.5, 'One Customer · One Journey · Every Channel', size=24, italic=True, color=MINT)
    text(s, 0.8, 5.55, 11, 0.4, 'Putra Irawan · untuk Kepala Divisi IT, PT Hasjrat Abadi · Oktober 2026',
         size=16, color='D8E2E5')
    text(s, 0.8, 6.0, 11.5, 0.4,
         'Usulan konsep assessment, bukan produk resmi PT Hasjrat Abadi. Prototype memakai data sintetis.',
         size=12, color='8DA1A8')
    notes(s, 'Pembuka: tugas ini menjawab permintaan management untuk mengubah WhatsApp dari channel komunikasi '
             'menjadi platform customer management yang terintegrasi. Saya membawa analisis, desain journey, '
             'arsitektur, dashboard penjualan dengan MS SQL, ide AI, dan prototype yang bisa didemokan.')


def slide_summary(prs, n):
    s = new_slide(prs)
    header(s, 'Ringkasan eksekutif', 'Pusatkan data pada customer, bukan pada chat atau HP sales')
    rect(s, 0.6, 1.7, 5.6, 1.45, fill=BG)
    text(s, 0.85, 1.88, 5.1, 0.3, 'PERMINTAAN MANAGEMENT', size=11, bold=True, color=MUTED)
    text(s, 0.85, 2.18, 5.1, 0.9, '“Mengubah WhatsApp dari sekadar communication channel menjadi customer '
                                  'management platform yang terintegrasi.”', size=15, italic=True)
    rect(s, 6.45, 1.7, 6.28, 1.45, fill=NAVY)
    text(s, 6.7, 1.88, 5.8, 0.3, 'JAWABAN', size=11, bold=True, color=MINT)
    text(s, 6.7, 2.18, 5.85, 0.9, 'Semua channel masuk ke satu profil customer, satu journey dengan SLA dan '
                                  'otomasi, tersinkron dengan DMS/ERP MS SQL, dipantau dashboard real-time.',
         size=15, color=WHITE)
    pillars = [
        ('One Customer', 'Nomor WhatsApp milik perusahaan, shared inbox, identity resolution lintas channel, '
                         'consent tersimpan.', 'Data di HP pribadi · customer mengulang informasi'),
        ('One Journey', 'Stage standar inquiry → SPK → DO → aftersales, routing otomatis, SLA 15 menit dengan '
                        'eskalasi, cadence follow-up.', '30% inquiry terlambat · tidak ada standar follow-up'),
        ('Every Channel, One Truth', 'Event dari MS SQL menghentikan prospecting saat membeli; reporting mart '
                                     '`rpt` untuk dashboard.', 'Management tidak melihat status · pembeli tetap '
                                                               'menerima prospecting · broadcast manual'),
    ]
    for i, (title, body, solves) in enumerate(pillars):
        x = 0.6 + i * 4.11
        rect(s, x, 3.45, 3.85, 3.3, fill=GREEN_LIGHT)
        badge(s, x + 0.25, 3.7, 0.5, str(i + 1), size=16)
        text(s, x + 0.9, 3.7, 2.8, 0.5, title, size=17, bold=True, anchor=MSO_ANCHOR.MIDDLE)
        text(s, x + 0.25, 4.4, 3.35, 1.3, body, size=14)
        text(s, x + 0.25, 5.75, 3.35, 0.3, 'MENYELESAIKAN', size=10, bold=True, color=GREEN_DARK)
        text(s, x + 0.25, 6.02, 3.35, 0.65, solves, size=12, color=MUTED)
    footer(s, n)
    notes(s, 'Satu kalimat jawaban: customer menjadi pusat data. Tiga pilar ini memetakan langsung ke tujuh '
             'gejala di skenario.')


def slide_outcomes(prs, n):
    s = new_slide(prs)
    header(s, 'Target outcome pilot 12 minggu', 'Keberhasilan diukur, bukan diasumsikan')
    tiles = [
        ('≥90%', 'inquiry mendapat respons manusia ≤15 menit', 'Baseline skenario: ±70%'),
        ('<5%', 'customer mengulang informasi saat pindah channel', 'Baseline skenario: ±20%'),
        ('100%', 'lead tercatat di CRM, 0 hanya di HP pribadi', 'Baseline: diukur saat discovery'),
        ('0', 'pembeli (SPK/DO) menerima pesan prospecting', 'Diblok otomatis saat kirim'),
        ('≤15 mnt', 'freshness dashboard prospect & penjualan', 'Baseline: manual / tidak ada'),
        ('↑', 'konversi lead → SPK dibanding baseline', 'Angka target ditetapkan setelah baseline'),
    ]
    for i, (value, label, base) in enumerate(tiles):
        col, row = i % 3, i // 3
        x, y = 0.6 + col * 4.11, 1.75 + row * 2.55
        rect(s, x, y, 3.85, 2.3, fill=WHITE, line=LINE)
        text(s, x + 0.3, y + 0.25, 3.3, 0.85, value, size=44, bold=True, color=GREEN)
        text(s, x + 0.3, y + 1.12, 3.3, 0.65, label, size=15, bold=True)
        text(s, x + 0.3, y + 1.8, 3.3, 0.3, base, size=12, color=MUTED)
    footer(s, n)
    notes(s, 'Target pilot. Baseline final diukur pada minggu 1–2 discovery; angka konversi target ditetapkan '
             'setelah baseline tersedia.')


def slide_scenario(prs, n):
    s = new_slide(prs)
    header(s, 'Task 1 · Problem diagnosis', 'Tujuh gejala, satu pola: WhatsApp dipakai sebagai chat pribadi')
    stats = [('10.000', 'inquiry WhatsApp per bulan', NAVY, WHITE, MINT),
             ('3.000', 'inquiry (30%) tidak di-follow-up tepat waktu', AMBER_LIGHT, INK, AMBER),
             ('2.000', 'inquiry (20%) mengulang informasi saat pindah channel', AMBER_LIGHT, INK, AMBER)]
    for i, (value, label, fill, color, vcolor) in enumerate(stats):
        x = 0.6 + i * 4.11
        rect(s, x, 1.7, 3.85, 2.05, fill=fill)
        text(s, x + 0.3, 1.88, 3.3, 0.95, value, size=54, bold=True, color=vcolor)
        text(s, x + 0.3, 2.9, 3.3, 0.7, label, size=15, color=color)
    others = ['Management tidak melihat status prospect secara real-time',
              'Sebagian sales menyimpan customer hanya di HP pribadi',
              'Tidak ada standard follow-up journey',
              'Broadcast marketing dilakukan manual',
              'Pembeli masih menerima prospecting yang tidak relevan']
    for i, label in enumerate(others):
        x = 0.6 + i * 2.47
        rect(s, x, 4.05, 2.3, 1.9, fill=BG)
        badge(s, x + 0.2, 4.25, 0.42, str(i + 3), fill=NAVY_SOFT, size=13)
        text(s, x + 0.2, 4.82, 1.95, 1.05, label, size=13)
    rect(s, 0.6, 6.17, 12.13, 0.6, fill=GREEN_LIGHT, radius=0.3)
    text(s, 0.85, 6.17, 11.7, 0.6, [[('Diagnosis: ', {'bold': True, 'color': GREEN_DARK}),
                                     ('tidak ada pemilik data, tidak ada proses terukur, dan tidak ada sambungan '
                                      'antara percakapan dengan status penjualan.', {})]],
         size=15, anchor=MSO_ANCHOR.MIDDLE)
    footer(s, n)
    notes(s, 'Ketujuh gejala bukan tujuh masalah terpisah. Semuanya muncul karena WhatsApp dipakai sebagai alat '
             'chat individu, bukan sistem perusahaan.')


def slide_root_causes(prs, n):
    s = new_slide(prs)
    header(s, 'Task 1 · Akar masalah', 'Enam akar masalah di balik tujuh gejala')
    causes = [
        ('A', 'Tidak ada system of record', 'Customer “milik” sales, bukan perusahaan; tidak ada kebijakan '
                                            'ownership data.', 'Governance · Technology', 'Gejala 3, 4'),
        ('B', 'Tidak ada lead management', 'Tanpa routing, SLA timer, eskalasi, dan pengaturan kapasitas; '
                                           'after-hours tidak tertangani.', 'Process · People', 'Gejala 1'),
        ('C', 'Identitas terfragmentasi', 'Setiap channel dan setiap HP adalah silo; tidak ada histori '
                                          'percakapan bersama.', 'Data · Technology', 'Gejala 2'),
        ('D', 'Tidak ada journey standar', 'Stage, cadence, dan exit criteria tidak didefinisikan; kualitas '
                                           'bergantung individu.', 'Process', 'Gejala 5'),
        ('E', 'Marketing lepas dari penjualan', 'Tidak ada consent, segmentasi, dan suppression berbasis status '
                                                'SPK/DO.', 'Process · Data · Governance', 'Gejala 6, 7'),
        ('F', 'Tidak ada lapisan KPI', 'Definisi metrik dan data segar tidak tersedia untuk management.',
         'Data', 'Gejala 3'),
    ]
    for i, (letter, title, body, cat, links) in enumerate(causes):
        col, row = i % 3, i // 3
        x, y = 0.6 + col * 4.11, 1.7 + row * 2.62
        rect(s, x, y, 3.85, 2.4, fill=GREEN_LIGHT if row == 0 else BG)
        badge(s, x + 0.25, y + 0.25, 0.48, letter, fill=GREEN if row == 0 else NAVY_SOFT, size=16)
        text(s, x + 0.9, y + 0.25, 2.8, 0.48, title, size=16, bold=True, anchor=MSO_ANCHOR.MIDDLE)
        text(s, x + 0.25, y + 0.9, 3.35, 0.9, body, size=13)
        text(s, x + 0.25, y + 1.9, 2.2, 0.3, cat, size=11, color=MUTED)
        chip(s, x + 2.55, y + 1.85, 1.1, 0.34, links, fill=WHITE, size=11, bold=True)
    footer(s, n)
    notes(s, 'Kategori People/Process/Technology/Data/Governance. Prioritas pertama adalah A dan B karena '
             'paling cepat memberi dampak.')


def slide_symptom_table(prs, n):
    s = new_slide(prs)
    header(s, 'Task 1 · Gejala → solusi', 'Setiap gejala punya solusi yang dapat diukur')
    rows = [['#', 'Gejala', 'Akar masalah', 'Solusi OneCRM'],
            ['1', '30% inquiry terlambat di-follow-up', 'Tanpa routing, SLA, eskalasi', 'Auto-ack <1 menit, routing, SLA 15 menit + eskalasi, bot after-hours'],
            ['2', '20% customer mengulang informasi', 'Identitas terfragmentasi', 'Identity resolution, timeline lintas channel, ringkasan saat handover'],
            ['3', 'Status prospect tidak terlihat', 'Tidak ada system of record & KPI', 'Pipeline stage standar, dashboard CRM + penjualan MS SQL'],
            ['4', 'Customer hanya di HP pribadi', 'Nomor & akun milik pribadi', 'Nomor WA perusahaan, shared inbox di PWA, bulk reassign'],
            ['5', 'Tidak ada standar follow-up', 'Tidak ada playbook & cadence', 'Journey 9 tahap, cadence H+0/1/3/7/14, task otomatis'],
            ['6', 'Broadcast manual', 'Tidak ada campaign tool & consent', 'Campaign studio: segmen, template disetujui, frequency cap'],
            ['7', 'Pembeli menerima prospecting', 'Marketing tidak terhubung ke SPK/DO', 'Event MS SQL → Order/Won → suppression + send-time check']]
    table(s, 0.6, 1.7, 12.13, rows, [0.5, 3.4, 3.4, 4.83], size=13, header_size=13, row_h=0.6)
    footer(s, n)
    notes(s, 'Tabel lengkap dengan kolom dampak bisnis ada di dokumen tertulis, bagian 2.2.')


def slide_impact(prs, n):
    s = new_slide(prs)
    header(s, 'Task 1 · Dampak & prioritas', 'Follow-up terlambat: ±30 SPK berpotensi hilang tiap bulan')
    steps = [('10.000', 'inquiry WhatsApp per bulan', 'skenario'),
             ('±4.000', 'lead unik  (÷ 2,5 inquiry per lead)', 'asumsi'),
             ('±1.200', 'lead terlambat  (× 30%)', 'skenario'),
             ('±30 SPK', 'hilang per bulan  (× selisih konversi 2,5 pp)', 'asumsi'),
             ('±Rp3,9 M', 'nilai penjualan/bulan  (× Rp130 juta/unit)', 'asumsi')]
    for i, (value, label, tag) in enumerate(steps):
        y = 1.7 + i * 0.86
        last = i == len(steps) - 1
        rect(s, 0.6, y, 6.6, 0.72, fill=NAVY if last else (GREEN_LIGHT if tag == 'skenario' else BG))
        text(s, 0.85, y, 1.75, 0.72, value, size=20, bold=True, color=MINT if last else GREEN_DARK,
             anchor=MSO_ANCHOR.MIDDLE)
        text(s, 2.65, y, 3.4, 0.72, label, size=13, color=WHITE if last else INK, anchor=MSO_ANCHOR.MIDDLE)
        chip(s, 6.0, y + 0.2, 1.0, 0.32, tag, fill=WHITE if not last else NAVY_SOFT,
             color=AMBER if tag == 'asumsi' else GREEN_DARK, size=10, bold=True)
    text(s, 0.6, 6.12, 6.6, 0.7, 'Asumsi untuk divalidasi saat discovery. Nilai penjualan, bukan margin. '
                                 'Tambahan ±100 jam kerja sales/bulan hilang untuk menanyakan ulang informasi.',
         size=12, color=MUTED)
    # priority matrix
    mx, my, cw, ch = 8.2, 2.05, 2.25, 2.0
    text(s, 7.65, 1.7, 5, 0.3, 'PRIORITAS: IMPACT × EFFORT', size=11, bold=True, color=MUTED)
    quads = [(0, 0, 'Quick wins', 'Nomor WA perusahaan + shared inbox, auto-ack, routing & SLA, suppression '
                                  'pembeli', GREEN, WHITE),
             (1, 0, 'Strategic', 'Identity resolution, event MS SQL, dashboard penjualan', NAVY, WHITE),
             (0, 1, 'Fill-ins', 'Template follow-up standar, opt-out keyword', GREEN_LIGHT, INK),
             (1, 1, 'Fase lanjut', 'AI ringkasan & scoring, IG/FB/TikTok penuh', BG, INK)]
    for col, row, title, body, fill, color in quads:
        x, y = mx + col * (cw + 0.12), my + row * (ch + 0.12)
        rect(s, x, y, cw, ch, fill=fill, radius=0.06)
        text(s, x + 0.18, y + 0.15, cw - 0.3, 0.35, title, size=14, bold=True, color=color)
        text(s, x + 0.18, y + 0.55, cw - 0.3, ch - 0.65, body, size=12, color=color)
    text(s, 7.65, my + 0.2, 0.45, ch, 'Impact tinggi', size=10, color=MUTED)
    text(s, 7.65, my + ch + 0.32, 0.45, ch, 'Impact sedang', size=10, color=MUTED)
    text(s, mx, my + 2 * ch + 0.3, cw, 0.3, 'Effort rendah', size=10, color=MUTED, align=PP_ALIGN.CENTER)
    text(s, mx + cw + 0.12, my + 2 * ch + 0.3, cw, 0.3, 'Effort tinggi', size=10, color=MUTED,
         align=PP_ALIGN.CENTER)
    footer(s, n)
    notes(s, 'Formula sengaja transparan. Semua asumsi (2,5 inquiry per lead, selisih konversi 2,5 poin, '
             'Rp130 juta per unit) dihitung ulang dengan baseline nyata pada discovery.')


def slide_journey_map(prs, n):
    s = new_slide(prs)
    header(s, 'Task 2 · Customer journey', 'Satu journey standar dari iklan sampai aftersales')
    row1 = [('0', 'Awareness', 'UTM & ad ID ditempel ke link dan form'),
            ('1', 'Inquiry & capture', 'Auto-ack <1 menit, consent, identity resolution'),
            ('2', 'Kualifikasi', 'Respons manusia ≤15 menit; bot after-hours; Hot/Warm/Cold'),
            ('3', 'Assignment & konsultasi', 'Routing per cabang, produk, kapasitas'),
            ('4', 'Test drive / kunjungan', 'Booking, reminder H-1 & 2 jam, no-show task')]
    row2 = [('5', 'Penawaran & negosiasi', 'Quotation, simulasi kredit, leasing, trade-in'),
            ('6', 'SPK (Order)', 'Dari MS SQL: stage Order, prospecting berhenti'),
            ('7', 'DO / serah terima', 'Stage Won, survey CSAT/NPS H+1'),
            ('8', 'Aftersales & loyalty', 'Servis, STNK/asuransi, referral, repeat')]
    text(s, 0.6, 1.62, 6, 0.3, 'PROSPECTING', size=11, bold=True, color=GREEN)
    cw = 2.25
    for i, (num, title, body) in enumerate(row1):
        x = 0.6 + i * (cw + 0.22)
        rect(s, x, 1.95, cw, 1.85, fill=GREEN_LIGHT)
        badge(s, x + 0.18, 2.12, 0.42, num, size=13)
        text(s, x + 0.7, 2.08, cw - 0.8, 0.52, title, size=14, bold=True, anchor=MSO_ANCHOR.MIDDLE)
        text(s, x + 0.18, 2.72, cw - 0.33, 1.0, body, size=12)
        if i < len(row1) - 1:
            arrow(s, x + cw + 0.02, 2.87, x + cw + 0.2, 2.87)
    text(s, 0.6, 3.98, 6, 0.3, 'ORDER, SERAH TERIMA & KEPEMILIKAN', size=11, bold=True, color=NAVY_SOFT)
    for i, (num, title, body) in enumerate(row2):
        x = 0.6 + i * (cw + 0.22)
        fill = NAVY if num in ('6', '7') else BG
        color = WHITE if fill == NAVY else INK
        rect(s, x, 4.31, cw, 1.6, fill=fill)
        badge(s, x + 0.18, 4.48, 0.42, num, fill=MINT if fill == NAVY else NAVY_SOFT,
              color=NAVY if fill == NAVY else WHITE, size=13)
        text(s, x + 0.7, 4.44, cw - 0.8, 0.52, title, size=14, bold=True, color=color, anchor=MSO_ANCHOR.MIDDLE)
        text(s, x + 0.18, 5.05, cw - 0.33, 0.8, body, size=12, color=color)
        if i < len(row2) - 1:
            arrow(s, x + cw + 0.02, 5.11, x + cw + 0.2, 5.11)
    x = 0.6 + 4 * (cw + 0.22)
    rect(s, x, 4.31, cw, 1.6, fill=AMBER_LIGHT)
    text(s, x + 0.18, 4.45, cw - 0.3, 0.3, 'EXIT PATHS', size=11, bold=True, color=AMBER)
    bullets(s, x + 0.18, 4.8, cw - 0.3, 1.1, ['Lost + reason code', 'Nurture', 'Opt-out', 'Purchased → suppress'],
            size=12, space_after=1, bullet_color=AMBER)
    rect(s, 0.6, 6.15, 12.13, 0.62, fill=WHITE, line=LINE, radius=0.3)
    text(s, 0.85, 6.15, 11.7, 0.62, [[('Prinsip: ', {'bold': True, 'color': GREEN_DARK}),
                                      ('satu profil, banyak opportunity · customer tidak mengulang informasi · '
                                       'status SPK/DO dari sistem sumber · komunikasi mengikuti tahap', {})]],
         size=14, anchor=MSO_ANCHOR.MIDDLE)
    footer(s, n)
    notes(s, 'Tahap 6 dan 7 digerakkan oleh event dari DMS/ERP MS SQL, bukan diketik ulang oleh sales. '
             'Itulah yang menghentikan prospecting ke pembeli.')


def slide_journey_table(prs, n):
    s = new_slide(prs)
    header(s, 'Task 2 · Detail tahap', 'Setiap tahap punya otomasi, pemilik, SLA, dan exit criteria')
    rows = [['Tahap', 'Otomasi utama', 'Owner & SLA', 'Exit criteria'],
            ['1 Inquiry & capture', 'Auto-ack, consent, E.164, merge profil, SLA timer', 'Sistem, <1 menit', 'Lead ter-assign'],
            ['2 Kualifikasi', 'Bot after-hours, scoring Hot/Warm/Cold', 'Sales ≤15 menit', 'Qualified / Nurture / Lost'],
            ['3 Konsultasi', 'Routing, template brosur & simulasi kredit', 'Sales; eskalasi menit 15', 'Jadwal test drive'],
            ['4 Test drive', 'Booking, reminder H-1 & 2 jam, no-show task', 'Sales + showroom', 'Hadir / reschedule'],
            ['5 Penawaran', 'Quotation, leasing, trade-in, reminder H+2', 'Sales; diskon oleh supervisor', 'SPK / Nurture / Lost'],
            ['6 SPK', 'erp.spk.created → Order, stop prospecting', 'Sales + admin + leasing', 'DO / SPK batal'],
            ['7 DO', 'erp.do.delivered → Won, survey H+1', 'Sales + delivery', 'Masuk aftersales'],
            ['8 Aftersales', 'Servis, STNK/asuransi, referral, repeat', 'Service advisor, CRM', 'Opportunity baru']]
    table(s, 0.6, 1.7, 12.13, rows, [2.4, 4.6, 2.85, 2.28], size=13, header_size=13, row_h=0.52)
    footer(s, n)
    notes(s, 'Data yang ditangkap per tahap (model, budget, cash/kredit, timeline, lokasi, CSAT) dirinci di '
             'dokumen tertulis bagian 3.2.')


def slide_cadence(prs, n):
    s = new_slide(prs)
    header(s, 'Task 2 · Cadence follow-up', 'Follow-up terstandar yang berhenti otomatis')
    points = [('H+0', '≤15 menit', 'Respons manusia pertama'), ('H+1', 'Info relevan', 'Brosur / simulasi'),
              ('H+3', 'Ajak bertemu', 'Test drive / showroom'), ('H+7', 'Penawaran', 'Tanyakan kendala'),
              ('H+14', 'Check-in akhir', 'Tawarkan nurture'), ('Bulanan', 's/d 90 hari', 'Konten bernilai')]
    y = 2.35
    line(s, 1.6, y, 11.73, y, color=GREEN, width=3)
    for i, (when, what, detail) in enumerate(points):
        x = 1.6 + i * 2.026
        badge(s, x - 0.22, y - 0.22, 0.44, '', fill=GREEN if i < 5 else MINT)
        text(s, x - 1.0, 1.65, 2.0, 0.4, when, size=18, bold=True, color=GREEN_DARK, align=PP_ALIGN.CENTER)
        text(s, x - 1.0, 2.75, 2.0, 0.35, what, size=14, bold=True, align=PP_ALIGN.CENTER)
        text(s, x - 1.0, 3.1, 2.0, 0.35, detail, size=12, color=MUTED, align=PP_ALIGN.CENTER)
    rect(s, 0.6, 3.85, 5.95, 2.3, fill=GREEN_LIGHT)
    text(s, 0.85, 4.05, 5.5, 0.35, 'Stop rules otomatis', size=16, bold=True)
    bullets(s, 0.85, 4.5, 5.5, 2.2, ['Customer membalas → kembali ke percakapan aktif',
                                     'SPK atau DO tercatat di MS SQL',
                                     'Opt-out (STOP / BERHENTI)',
                                     'Status Lost dengan reason code',
                                     'Frequency cap: maks. 2 pesan marketing per minggu'], size=14, space_after=3)
    rect(s, 6.78, 3.85, 5.95, 2.3, fill=BG)
    text(s, 7.03, 4.05, 5.5, 0.35, 'Kontinuitas omnichannel', size=16, bold=True)
    bullets(s, 7.03, 4.5, 5.5, 2.2, ['Satu profil: WA, IG, Messenger, TikTok, website, walk-in',
                                     'Ringkasan percakapan saat pindah channel atau sales',
                                     'Timeline: interaksi, stage, consent, SPK/DO, servis',
                                     'Pesan transaksional tetap jalan setelah pembelian'], size=14, space_after=3)
    footer(s, n)
    notes(s, 'Cadence dapat dikonfigurasi per cabang/produk. Respons manusia dihitung terpisah dari auto-ack.')


def slide_architecture(prs, n):
    s = new_slide(prs)
    header(s, 'Task 3 · High-level architecture', 'Arsitektur berlapis, event-driven, ERP tetap sumber kebenaran')
    lanes = [('Channels', ['WhatsApp\nCloud API · BSP', 'Instagram\nDM · Lead Ads', 'Facebook\nMessenger · Lead Ads',
                           'TikTok\nLead Generation', 'Website\nform · chat · UTM', 'Walk-in / telepon\ninput PWA'],
              WHITE, LINE),
             ('Integration', ['Webhook gateway\nverifikasi signature', 'Queue / event bus\nretry · DLQ',
                              'Channel adapters\nidempotency key'], WHITE, LINE),
             ('Core services · Go', ['Identity\nresolution', 'Lead mgmt\nrouting · SLA', 'Shared\ninbox',
                                     'Journey\nengine', 'Consent\npreferensi', 'Campaign\nsuppression',
                                     'Task &\neskalasi'], GREEN_LIGHT, None),
             ('Data', ['OneCRM DB\nMySQL demo / MS SQL', 'Audit & event log', 'DMS/ERP MS SQL\nread-only · CT/CDC',
                       'Reporting mart\nMS SQL schema rpt'], NAVY, None)]
    lane_y = [1.7, 2.92, 3.98, 5.2]
    lane_h = [0.95, 0.8, 0.95, 0.95]
    left, right = 1.95, 9.95
    for (label, items, fill, line_color), y, h in zip(lanes, lane_y, lane_h):
        text(s, 0.6, y, 1.3, h, label, size=13, bold=True, color=GREEN_DARK, anchor=MSO_ANCHOR.MIDDLE)
        gap = 0.08
        w = (right - left - gap * (len(items) - 1)) / len(items)
        for i, item in enumerate(items):
            chip(s, left + i * (w + gap), y, w, h, item, fill=fill, line=line_color, size=12 if len(items) < 7 else 11,
                 bold=True, color=WHITE if fill == NAVY else INK)
    for y in (2.68, 3.75, 4.96):
        arrow(s, 5.95, y, 5.95, y + 0.22)
    # right column
    blocks = [('Experience', 'Sales PWA: inbox, task, pipeline\nDashboard supervisor & mgmt\nCampaign studio',
               GREEN_LIGHT, INK),
              ('AI gateway', 'Server-side, model-agnostic\nKonteks minimal · review manusia', AMBER_LIGHT, INK),
              ('Security & governance', 'SSO Entra ID · RBAC per peran\nAudit log · masking PII · enkripsi\n'
                                        'Consent & retensi sesuai UU PDP', NAVY, WHITE)]
    y = 1.7
    for title, body, fill, color in blocks:
        h = 1.42 if title != 'AI gateway' else 1.05
        rect(s, 10.15, y, 2.58, h, fill=fill)
        text(s, 10.33, y + 0.12, 2.25, 0.3, title, size=13, bold=True, color=color)
        text(s, 10.33, y + 0.45, 2.3, h - 0.5, body.split('\n'), size=11, color=color)
        y += h + 0.15
    text(s, 0.6, 6.32, 12.13, 0.5, 'Prinsip: nomor & akun milik perusahaan · event idempoten · ERP MS SQL '
                                   'adalah sumber SPK/DO/stok · otomasi dapat dikonfigurasi tanpa deploy',
         size=13, color=MUTED)
    footer(s, n)
    notes(s, 'Prototype mengimplementasikan sebagian: Next.js + TypeScript, Go API, MySQL untuk CRM, '
             'MS SQL Server untuk reporting penjualan, dan DeepSeek melalui backend.')


def slide_omnichannel(prs, n):
    s = new_slide(prs)
    header(s, 'Task 3 · Otomasi lead omnichannel', 'Lima channel digital + walk-in, satu alur pemrosesan lead')
    sources = [('WhatsApp', 'Webhook Cloud API; referral ad ID'), ('Website', 'Form + CAPTCHA, consent, UTM'),
               ('Instagram', 'DM API; Lead Ads webhook leadgen'), ('Facebook', 'Messenger; Lead Ads leadgen'),
               ('TikTok', 'Lead Generation (Instant Form)'), ('Walk-in', 'Input cepat di PWA sales')]
    for i, (name, how) in enumerate(sources):
        y = 1.72 + i * 0.8
        rect(s, 0.6, y, 3.4, 0.68, fill=WHITE, line=LINE)
        text(s, 0.8, y + 0.06, 3.1, 0.3, name, size=14, bold=True)
        text(s, 0.8, y + 0.36, 3.1, 0.3, how, size=11, color=MUTED)
        arrow(s, 4.05, y + 0.34, 4.6, y + 0.34, color='B8C4C7', width=1.25)
    steps = [('Capture', 'event + idempotency key'), ('Normalisasi', 'nomor → E.164, email, social ID'),
             ('Identity & dedupe', 'cari / merge golden customer ID'), ('Consent', 'catat opt-in & preferensi'),
             ('Scoring', 'intent, kelengkapan, perilaku'), ('Routing', 'cabang · produk · kapasitas'),
             ('SLA timer & task', 'notifikasi di PWA sales')]
    for i, (title, body) in enumerate(steps):
        y = 1.72 + i * 0.67
        badge(s, 4.85, y + 0.06, 0.44, str(i + 1), size=13)
        text(s, 5.45, y, 2.6, 0.56, title, size=14, bold=True, anchor=MSO_ANCHOR.MIDDLE)
        text(s, 7.75, y, 2.5, 0.56, body, size=12, color=MUTED, anchor=MSO_ANCHOR.MIDDLE)
    rect(s, 10.35, 1.72, 2.38, 4.68, fill=NAVY)
    text(s, 10.55, 1.92, 2.0, 0.35, 'HASIL', size=11, bold=True, color=MINT)
    bullets(s, 10.55, 2.32, 2.05, 3.9, ['Satu lead per customer, bukan per channel', 'Owner jelas dalam <1 menit',
                                        'SLA 15 menit berjalan', 'Sumber lead tersimpan untuk attribution SPK'],
            size=13, color=WHITE, space_after=8, bullet_color=MINT)
    text(s, 0.6, 6.5, 12.13, 0.35, 'Akses API Instagram, Facebook, dan TikTok mengikuti app review dan entitlement '
                                   'tiap platform; pilot dimulai dari WhatsApp + website.', size=12, color=MUTED)
    footer(s, n)
    notes(s, 'Replay event yang sama tidak membuat lead ganda — ini didemonstrasikan di Integration Hub prototype.')


def slide_automation(prs, n):
    s = new_slide(prs)
    header(s, 'Task 3 · Katalog otomasi', 'Dua belas aturan otomasi menggantikan kerja manual')
    rows = [['Trigger', 'Aksi', 'KPI'],
            ['inquiry.created', 'Auto-ack + consent, merge profil, routing, mulai SLA', '% auto-ack <1 menit'],
            ['sla.warning (menit 10)', 'Push notification ke owner', 'First response time'],
            ['sla.breached (menit 15)', 'Eskalasi supervisor, reassign', '% respons ≤15 menit'],
            ['inquiry.after_hours', 'Bot kualifikasi, janji waktu respons', '% lead terkualifikasi'],
            ['lead.scored = Hot', 'Task prioritas + link booking test drive', 'Hot → test drive'],
            ['testdrive.booked', 'Reminder H-1 & 2 jam; no-show → task', 'Show rate'],
            ['followup.due H+1/3/7', 'Task + template sesuai tahap', '% follow-up lengkap'],
            ['erp.spk.created', 'Stage Order, hentikan prospecting', '0 pembeli diprospek'],
            ['erp.do.delivered', 'Stage Won, survey H+1, aftersales journey', 'CSAT / NPS'],
            ['optout.received', 'Cabut consent, blok marketing', 'Diproses <1 menit'],
            ['user.deactivated', 'Bulk reassign lead & percakapan', '0 lead tanpa owner'],
            ['campaign.send', 'Cek ERP segar, consent, stage, cap, template', '0 pelanggaran']]
    table(s, 0.6, 1.62, 12.13, rows, [3.0, 6.0, 3.13], size=12, header_size=12, row_h=0.39)
    footer(s, n)
    notes(s, 'Event erp.spk.created dan erp.do.delivered dibaca dari MS SQL melalui Change Tracking/CDC.')


def slide_dashboard_idea(prs, n):
    s = new_slide(prs)
    header(s, 'Dashboard penjualan · Ide', 'Satu dashboard untuk direksi, cabang, supervisor, dan marketing')
    users = [('Direksi', 'On-track terhadap target? Channel mana yang menghasilkan penjualan?'),
             ('Kepala cabang', 'SPK & DO vs target, backlog SPK, sales yang perlu dibantu'),
             ('Sales supervisor', 'Leaderboard dan konversi lead → SPK per sumber'),
             ('Marketing', 'Lead per sumber dan konversinya ke SPK/DO')]
    for i, (role, question) in enumerate(users):
        y = 1.7 + i * 1.27
        rect(s, 0.6, y, 4.9, 1.12, fill=GREEN_LIGHT)
        text(s, 0.85, y + 0.13, 4.4, 0.32, role, size=15, bold=True)
        text(s, 0.85, y + 0.47, 4.45, 0.62, question, size=13)
    rows = [['KPI', 'Definisi (identik di TypeScript & SQL)'],
            ['SPK (gross)', 'SPK dibuat dalam periode, semua status'],
            ['SPK batal', 'SPK dalam periode berstatus Cancelled'],
            ['DO & revenue', 'Unit & nilai dengan tanggal DO dalam periode'],
            ['Target · on-pace', 'Target bulanan penuh dari bulan yang disentuh; pro-rata hari'],
            ['Outstanding SPK', 'Backlog akhir periode: tidak batal, belum DO'],
            ['Kredit share', 'DO melalui leasing ÷ total DO'],
            ['Lead → SPK', 'SPK ÷ lead unik CRM, per sumber lead'],
            ['Trend 6 bulan', 'SPK, DO, target per bulan']]
    table(s, 5.75, 1.7, 6.98, rows, [2.0, 4.98], size=12, header_size=12, row_h=0.535)
    text(s, 5.75, 6.65, 6.98, 0.3, 'Dashboard hanya membawa angka hitungan; rasio dihitung di UI agar tidak ada '
                                   'selisih pembulatan.', size=11, color=MUTED)
    footer(s, n)
    notes(s, 'Definisi KPI ini diimplementasikan sama persis di lib/sales.ts dan di prosedur '
             'rpt.usp_SalesDashboard, lalu diuji dengan golden test.')


def slide_dashboard_arch(prs, n):
    s = new_slide(prs)
    header(s, 'Dashboard penjualan · Integrasi MS SQL', 'Dari DMS/ERP MS SQL ke dashboard tanpa membebani OLTP')
    row1 = [('DMS/ERP MS SQL', 'OLTP: SPK, DO, target', NAVY, WHITE),
            ('Change Tracking / CDC', 'read-only login atau readable secondary', WHITE, INK),
            ('Go sync worker', 'incremental · rpt.EtlWatermark', WHITE, INK),
            ('Reporting mart rpt', 'onecrm_reporting · star schema', GREEN, WHITE)]
    row2 = [('rpt.usp_SalesDashboard', '@DateFrom, @DateTo, @Branch → 6 result set', GREEN_LIGHT, INK),
            ('Go API', 'GET /api/sales/dashboard · validasi filter', WHITE, INK),
            ('Next.js UI', 'label sumber & freshness', NAVY, WHITE)]
    bw, gap = 2.78, 0.33
    for i, (title, body, fill, color) in enumerate(row1):
        x = 0.6 + i * (bw + gap)
        rect(s, x, 1.75, bw, 1.25, fill=fill, line=LINE if fill == WHITE else None)
        text(s, x + 0.2, 1.88, bw - 0.4, 0.35, title, size=14, bold=True, color=color)
        text(s, x + 0.2, 2.27, bw - 0.4, 0.65, body, size=12, color=color)
        if i < len(row1) - 1:
            arrow(s, x + bw + 0.04, 2.37, x + bw + gap - 0.04, 2.37)
    arrow(s, 0.6 + 3 * (bw + gap) + bw / 2, 3.04, 0.6 + 3 * (bw + gap) + bw / 2, 3.4)
    for i, (title, body, fill, color) in enumerate(row2):
        x = 0.6 + (3 - i) * (bw + gap)
        rect(s, x, 3.45, bw, 1.25, fill=fill, line=LINE if fill == WHITE else None)
        text(s, x + 0.2, 3.58, bw - 0.4, 0.35, title, size=14, bold=True, color=color)
        text(s, x + 0.2, 3.97, bw - 0.4, 0.65, body, size=12, color=color)
        if i < len(row2) - 1:
            arrow(s, x - 0.04, 4.07, x - gap + 0.04, 4.07)
    x0 = 0.6
    rect(s, x0, 3.45, bw, 1.25, fill=BG)
    text(s, x0 + 0.2, 3.58, bw - 0.4, 0.35, 'OneCRM → FactLeadDaily', size=14, bold=True)
    text(s, x0 + 0.2, 3.97, bw - 0.4, 0.65, 'lead harian per sumber, dimuat sync worker untuk attribution', size=12)
    text(s, 0.6, 4.95, 6, 0.3, 'STAR SCHEMA rpt', size=11, bold=True, color=MUTED)
    for i, t in enumerate(['DimBranch', 'DimProduct', 'DimSalesperson', 'DimLeadSource', 'FactSpk',
                           'FactLeadDaily', 'FactTarget', 'EtlWatermark']):
        chip(s, 0.6 + i * 1.53, 5.27, 1.42, 0.38, t, fill=GREEN_LIGHT if t.startswith('Fact') else BG, size=11,
             bold=True)
    guards = [('Least privilege', 'EXECUTE pada prosedur saja'), ('RLS per cabang', 'SESSION_CONTEXT(branch)'),
              ('Koneksi aman', 'TLS, secret di server'), ('Refresh', '±5 menit + rekonsiliasi malam'),
              ('Reuse', 'Power BI memakai schema rpt')]
    for i, (title, body) in enumerate(guards):
        x = 0.6 + i * 2.47
        rect(s, x, 5.85, 2.3, 0.95, fill=WHITE, line=LINE)
        text(s, x + 0.18, 5.95, 2.0, 0.3, title, size=13, bold=True, color=GREEN_DARK)
        text(s, x + 0.18, 6.27, 2.0, 0.5, body, size=11, color=MUTED)
    footer(s, n)
    notes(s, 'Di prototype, skrip T-SQL ada di backend/sql/mssql/. Prosedur dijalankan pada SQL Server 2022 '
             'dan hasilnya dibandingkan dengan perhitungan TypeScript melalui golden test.')


def slide_dashboard_ui(prs, n):
    s = new_slide(prs)
    header(s, 'Dashboard penjualan · UI', 'Satu layar: on-track, sumber penjualan, sales, dan backlog')
    picture(s, 'sales-dashboard', 0.6, 1.65, 8.4, 5.2, 'Sales Dashboard')
    items = ['Filter periode (MTD, 3 bulan, 6 bulan) & cabang', 'Label sumber: MS SQL atau simulasi + freshness',
             'KPI: SPK, DO, revenue, achievement & on-pace', 'Trend 6 bulan: SPK vs DO vs target',
             'Target vs aktual per cabang', 'Product mix & kredit share', 'Lead-source funnel: lead → SPK → DO',
             'Leaderboard sales', 'Panel alur data & integrasi']
    for i, label in enumerate(items):
        y = 1.68 + i * 0.57
        badge(s, 9.3, y + 0.04, 0.36, str(i + 1), size=11)
        text(s, 9.8, y, 2.93, 0.5, label, size=13, anchor=MSO_ANCHOR.MIDDLE)
    footer(s, n)
    notes(s, 'Data di layar adalah data sintetis. Mode MS SQL menampilkan label sumber dan waktu sinkron terakhir.')


def slide_two_shots(prs, n, eyebrow, title, left, right, note):
    s = new_slide(prs)
    header(s, eyebrow, title)
    for i, (name, label, caption) in enumerate((left, right)):
        x = 0.6 + i * 6.22
        picture(s, name, x, 1.65, 5.91, 4.35, label)
        text(s, x, 6.12, 5.91, 0.3, label, size=14, bold=True)
        text(s, x, 6.45, 5.91, 0.45, caption, size=12, color=MUTED)
    footer(s, n)
    notes(s, note)


def slide_ai(prs, n):
    s = new_slide(prs)
    header(s, 'Penggunaan AI', 'AI sebagai asisten sales, dengan manusia tetap memutuskan')
    rect(s, 0.6, 1.7, 4.3, 5.05, fill=NAVY)
    text(s, 0.85, 1.9, 3.9, 0.3, 'SUDAH DITERAPKAN DI PROTOTYPE', size=11, bold=True, color=MINT)
    text(s, 0.85, 2.25, 3.9, 0.7, 'Draft follow-up dengan DeepSeek melalui Go backend', size=17, bold=True,
         color=WHITE)
    bullets(s, 0.85, 3.12, 3.85, 3.5, ['API key hanya di server', 'Hanya label produk & intent yang keluar — '
                                                                 'tanpa nama, nomor, atau chat mentah',
                                       'Tidak boleh mengarang harga, stok, diskon, kredit',
                                       'Wajib ditinjau sales sebelum dikirim',
                                       'Fallback template berlabel bila tanpa key'],
            size=14, color=WHITE, space_after=6, bullet_color=MINT)
    text(s, 0.85, 6.0, 3.85, 0.6, 'Coba di prototype: Omnichannel Inbox → “Draft dengan DeepSeek”', size=12,
         color=MINT)
    phases = [('Pilot', GREEN, ['Ekstraksi intent & entitas dari chat', 'Ringkasan percakapan saat handover',
                                'Draft balasan dengan RAG (price list, promo, stok)']),
              ('Fase 2', NAVY_SOFT, ['Lead scoring (rules → ML histori SPK)', 'Bot kualifikasi after-hours',
                                     'Next-best-action & waktu follow-up', 'Deteksi sentimen & komplain',
                                     'Segmentasi & variasi copy campaign']),
              ('Fase 3', AMBER, ['Forecast SPK/DO & stok', '“Tanya data” NL → SQL (view rpt)',
                                 'OCR KTP/NPWP/dokumen kredit', 'Voice note → CRM', 'Deteksi anomali'])]
    for i, (phase, color, items) in enumerate(phases):
        x = 5.15 + i * 2.56
        text(s, x, 1.72, 2.4, 0.3, phase.upper(), size=12, bold=True, color=color)
        for j, item in enumerate(items):
            y = 2.1 + j * 0.93
            rect(s, x, y, 2.4, 0.8, fill=GREEN_LIGHT if i == 0 else BG)
            text(s, x + 0.15, y, 2.12, 0.8, item, size=12, anchor=MSO_ANCHOR.MIDDLE)
    footer(s, n)
    notes(s, 'Portofolio lengkap (masalah, input, output, guardrail, KPI) ada di dokumen tertulis bagian 6.2.')


def slide_ai_governance(prs, n):
    s = new_slide(prs)
    header(s, 'Penggunaan AI · Tata kelola & pengalaman', 'Guardrail dulu, baru skala')
    rules = [('Human-in-the-loop', 'AI menyiapkan, manusia memutuskan dan mengirim'),
             ('Minimisasi PII', 'Masking nama/nomor; konteks seperlunya'),
             ('Tanpa janji komersial', 'Harga, stok, diskon, kredit hanya dari sistem sumber'),
             ('Evaluation set', 'Diuji setiap ganti model atau prompt'),
             ('Logging & biaya', 'Audit tanpa PII, kuota per fitur'),
             ('Model-agnostic', 'Gateway bisa berganti penyedia atau model lokal'),
             ('UU PDP', 'Dasar pemrosesan, consent, retensi, penilaian risiko')]
    for i, (title, body) in enumerate(rules):
        y = 1.7 + i * 0.72
        badge(s, 0.6, y + 0.08, 0.42, '✓', size=13)
        text(s, 1.2, y, 2.45, 0.6, title, size=14, bold=True, anchor=MSO_ANCHOR.MIDDLE)
        text(s, 3.65, y, 3.1, 0.6, body, size=12, color=MUTED, anchor=MSO_ANCHOR.MIDDLE)
    rect(s, 7.05, 1.7, 5.68, 5.05, fill=AMBER_LIGHT, line=AMBER, dash=True, radius=0.04)
    text(s, 7.3, 1.88, 5.2, 0.3, 'PENGALAMAN PENERAPAN AI SAYA', size=11, bold=True, color=AMBER)
    text(s, 7.3, 2.2, 5.2, 0.4, '[Diisi oleh Putra sebelum presentasi]', size=16, bold=True)
    prompts = ['[Proyek AI paling relevan: konteks, masalah, pengguna]',
               '[Solusi & tools/model yang dipakai, peran saya]',
               '[Hasil terukur: waktu, biaya, konversi, akurasi]',
               '[Pelajaran yang dibawa ke OneCRM]']
    bullets(s, 7.3, 2.8, 5.2, 3.0, prompts, size=14, space_after=10, bullet_color=AMBER)
    text(s, 7.3, 6.05, 5.2, 0.55, 'Template tabel lengkap tersedia di dokumen tertulis bagian 6.4.', size=12,
         color=MUTED)
    footer(s, n)
    notes(s, 'Bagian kanan sengaja dikosongkan: isi dengan pengalaman nyata sebelum presentasi.')


def slide_roadmap(prs, n):
    s = new_slide(prs)
    header(s, 'Roadmap', 'Pilot 12 minggu dengan gate yang jelas sebelum rollout')
    phases = [('Minggu 1–2', 'Discovery & baseline', ['Audit channel & nomor WA', 'Baseline SLA & konversi',
                                                       'Kontrak data MS SQL']),
              ('Minggu 3–5', 'Core CRM pilot', ['WA + website, 1–2 cabang', 'Shared inbox, routing, SLA',
                                                 'Customer 360°']),
              ('Minggu 6–8', 'ERP, lifecycle & dashboard', ['Change Tracking & sync', 'Suppression pembeli',
                                                             'Dashboard penjualan']),
              ('Minggu 9–12', 'Ekspansi & evaluasi', ['IG, FB, TikTok', 'Pilot AI ringkasan & draft',
                                                       'Evaluasi KPI, rencana rollout'])]
    weeks = [2, 3, 3, 4]
    unit = 12.13 / 12
    x = 0.6
    for (when, title, items), wk in zip(phases, weeks):
        w = wk * unit - 0.1
        rect(s, x, 1.75, w, 0.55, fill=GREEN if when != 'Minggu 1–2' else NAVY, radius=0.3)
        text(s, x + 0.15, 1.75, w - 0.3, 0.55, when, size=14, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
        rect(s, x, 2.45, w, 2.45, fill=GREEN_LIGHT)
        text(s, x + 0.18, 2.6, w - 0.36, 0.6, title, size=15, bold=True)
        bullets(s, x + 0.18, 3.25, w - 0.32, 1.6, items, size=12, space_after=3)
        x += wk * unit
    text(s, 0.6, 5.12, 6, 0.3, 'PILOT GATES', size=11, bold=True, color=MUTED)
    gates = ['Replay event tidak menggandakan lead', 'Pembeli & opt-out tidak menerima prospecting',
             'Audit SLA valid (auto-ack bukan respons)', 'Dashboard cocok dengan MS SQL',
             '≥90% lead pilot ditangani di OneCRM']
    for i, gate in enumerate(gates):
        x = 0.6 + i * 2.47
        rect(s, x, 5.45, 2.3, 1.3, fill=WHITE, line=LINE)
        badge(s, x + 0.18, 5.6, 0.36, '✓', size=11)
        text(s, x + 0.18, 6.05, 2.0, 0.65, gate, size=12)
    footer(s, n)
    notes(s, 'Estimasi bergantung pada akses API channel, kesiapan data ERP, dan kapasitas tim.')


def slide_risks(prs, n):
    s = new_slide(prs)
    header(s, 'Risiko & mitigasi', 'Risiko terbesar adalah adopsi, bukan teknologi')
    rows = [['Risiko', 'Mitigasi'],
            ['Sales enggan meninggalkan WA pribadi', 'Libatkan sales sejak desain, PWA lebih cepat, insentif berbasis data CRM, dukungan kepala cabang'],
            ['Akses API IG/FB/TikTok terbatas', 'Mulai WA + website; ikuti app review; Lead Ads sebagai jalur awal'],
            ['Kualitas data DMS/ERP tidak konsisten', 'Kontrak data, rekonsiliasi harian, label freshness, data steward'],
            ['Nomor WA terkena pembatasan', 'Template disetujui, opt-in, frequency cap, pantau quality rating'],
            ['Kepatuhan UU PDP', 'Consent tercatat, RBAC, masking, retensi, DPIA'],
            ['Biaya & kualitas AI', 'Use case berisiko rendah dulu, evaluation set, kuota biaya, fallback manual']]
    table(s, 0.6, 1.7, 12.13, rows, [4.0, 8.13], size=14, header_size=13, row_h=0.68)
    footer(s, n)
    notes(s, 'RACI ringan dan detail tim ada di dokumen tertulis bagian 7.4.')


def slide_close(prs, n):
    s = new_slide(prs, NAVY)
    text(s, 0.8, 0.9, 11, 0.4, 'LANGKAH BERIKUTNYA', size=13, bold=True, color=MINT)
    text(s, 0.8, 1.35, 11.5, 0.9, 'Mulai dari yang paling cepat memberi dampak', size=40, bold=True, color=WHITE)
    steps = [('1', 'Discovery 2 minggu', 'Ukur baseline SLA, konversi, dan inventaris nomor WA; sepakati KPI.'),
             ('2', 'Pilot WA + website', 'Nomor perusahaan, shared inbox, routing & SLA di 1–2 cabang.'),
             ('3', 'Sambungkan MS SQL', 'Suppression pembeli dan dashboard penjualan dari reporting mart.')]
    for i, (num, title, body) in enumerate(steps):
        x = 0.8 + i * 4.0
        rect(s, x, 2.75, 3.75, 2.3, fill=NAVY_SOFT)
        badge(s, x + 0.25, 3.0, 0.5, num, fill=MINT, color=NAVY, size=16)
        text(s, x + 0.25, 3.68, 3.25, 0.4, title, size=18, bold=True, color=WHITE)
        text(s, x + 0.25, 4.12, 3.3, 0.85, body, size=14, color='D8E2E5')
    text(s, 0.8, 5.55, 11.5, 0.5, 'Terima kasih', size=28, bold=True, color=WHITE)
    text(s, 0.8, 6.1, 11.5, 0.4, 'Prototype interaktif, source code, dan dokumen tertulis terlampir.',
         size=15, color=MINT)
    footer(s, n, dark=True)
    notes(s, 'Tutup dengan tiga langkah konkret dan ajak diskusi.')


def build():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    prs.core_properties.title = 'Hasjrat OneCRM — Assessment Putra Irawan'
    prs.core_properties.author = 'Putra Irawan'
    prs.core_properties.subject = 'Customer management platform terintegrasi (konsep assessment)'
    slide_title(prs)
    builders = [
        slide_summary, slide_outcomes, slide_scenario, slide_root_causes, slide_symptom_table, slide_impact,
        slide_journey_map, slide_journey_table, slide_cadence, slide_architecture, slide_omnichannel,
        slide_automation, slide_dashboard_idea, slide_dashboard_arch, slide_dashboard_ui,
        lambda p, n: slide_two_shots(
            p, n, 'Prototype · UI', 'CRM Overview dan Omnichannel Inbox dalam satu workspace',
            ('crm-overview', 'CRM Overview', 'Inquiry vs lead unik, pipeline, konversi, SLA respons manusia'),
            ('inbox', 'Omnichannel Inbox', 'Shared inbox lintas channel dengan draft AI yang wajib ditinjau'),
            'Demo: filter cabang, buka lead yang melewati SLA, buat draft, catat balasan.'),
        lambda p, n: slide_two_shots(
            p, n, 'Prototype · UI', 'Journey Automation dan AI Playbook',
            ('journey', 'Journey Automation', 'Journey standar dengan aturan routing, SLA, dan suppression'),
            ('ai-playbook', 'AI Playbook', 'AI yang sudah diterapkan, portofolio usulan, dan governance'),
            'Demo: tunjukkan aturan otomasi dan bagaimana AI dibatasi guardrail.'),
        slide_ai, slide_ai_governance, slide_roadmap, slide_risks, slide_close,
    ]
    for i, builder in enumerate(builders, start=2):
        builder(prs, i)
    assert len(prs.slides) == TOTAL, len(prs.slides)
    prs.save(OUT)
    print(f'Wrote {OUT} ({len(prs.slides)} slides)')


if __name__ == '__main__':
    build()
